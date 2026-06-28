#!/usr/bin/env python3
import json
import sys
from openpyxl import load_workbook

SHEETS = [
    {'name': 'Sheet1', 'batches': ['25-26', '22-23', '19-20', '16-17']},
    {'name': 'Sheet2', 'batches': ['24-25', '21-22', '18-19']},
    {'name': 'Sheet3', 'batches': ['23-24', '20-21', '17-18']},
]


class PeopleStore:
    def __init__(self):
        self._people = {}
        self._next_id = 1

    def upsert(self, name, batch_year):
        key = (name, batch_year)
        person = self._people.get(key)
        if person is None:
            person = {
                'id': self._next_id,
                'name': name,
                'batchYear': batch_year,
                'mentorId': None,
            }
            self._people[key] = person
            self._next_id += 1
        return person

    def to_list(self):
        return list(self._people.values())


def is_empty_row(row):
    return all(cell.value is None for cell in row)


def cell_value(cell):
    if cell.value is None:
        return None
    val = str(cell.value).strip()
    return val if val else None


def find_separator(rows):
    for i in range(len(rows) - 2):
        if all(is_empty_row(rows[i + j]) for j in range(3)):
            return i
    return None


def process_chain_rows(rows, store, batches):
    num_cols = len(batches)
    for row in rows:
        if is_empty_row(row):
            continue

        last_person = None
        for col_idx in range(num_cols - 1, -1, -1):
            name = cell_value(row[col_idx])
            if name is None:
                continue

            person = store.upsert(name, batches[col_idx])
            if last_person is not None:
                person['mentorId'] = last_person['id']
            last_person = person


def main():
    if len(sys.argv) < 2:
        print('Usage: python parse.py <input.xlsx> [output.json]')
        sys.exit(1)

    input_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else 'data.json'

    all_batches = []
    for sheet_cfg in SHEETS:
        all_batches.extend(sheet_cfg['batches'])

    if len(all_batches) != len(set(all_batches)):
        print('Error: duplicate batch years found across sheets')
        sys.exit(1)

    wb = load_workbook(input_path, data_only=True)
    store = PeopleStore()

    for sheet_cfg in SHEETS:
        sheet_name = sheet_cfg['name']
        batches = sheet_cfg['batches']
        num_cols = len(batches)

        if sheet_name not in wb.sheetnames:
            print(f'Warning: sheet "{sheet_name}" not found, skipping')
            continue

        ws = wb[sheet_name]
        rows = list(ws.iter_rows(min_row=3, max_col=num_cols, values_only=False))
        sep_idx = find_separator(rows)

        if sep_idx is None:
            print(f'Warning ({sheet_name}): no 3-empty-row separator. '
                  'Treating all data as chain rows.')
            process_chain_rows(rows, store, batches)
        else:
            process_chain_rows(rows[:sep_idx], store, batches)
            process_chain_rows(rows[sep_idx + 3:], store, batches)

        print(f'  {sheet_name}: processed ({len(batches)} cols)')

    result = {'jasers': store.to_list()}

    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(result, f, indent=2, ensure_ascii=False)

    print(f'Done — {len(result["jasers"])} JASERs -> {output_path}')


if __name__ == '__main__':
    main()
